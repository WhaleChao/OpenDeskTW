from __future__ import annotations

import base64
from datetime import datetime, timedelta, timezone
import importlib.util
from pathlib import Path
import shutil
import tempfile

import fitz
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import rsa
from cryptography.hazmat.primitives.serialization import pkcs12
from cryptography.x509.oid import NameOID
from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
CORE_PATH = ROOT / "src-tauri/resources/acropdf-core/embedded_core.py"
SPEC = importlib.util.spec_from_file_location("document_pdf_core", CORE_PATH)
assert SPEC and SPEC.loader
core = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(core)


def make_image(path: Path, color: int) -> None:
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 64, 48), False)
    pixmap.clear_with(color)
    pixmap.save(path)


def make_certificate(path: Path, password: bytes) -> None:
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "全能文件工作台測試")])
    now = datetime.now(timezone.utc)
    certificate = (
        x509.CertificateBuilder()
        .subject_name(name)
        .issuer_name(name)
        .public_key(key.public_key())
        .serial_number(x509.random_serial_number())
        .not_valid_before(now - timedelta(minutes=1))
        .not_valid_after(now + timedelta(days=1))
        .sign(key, hashes.SHA256())
    )
    path.write_bytes(
        pkcs12.serialize_key_and_certificates(
            b"document-workbench",
            key,
            certificate,
            None,
            serialization.BestAvailableEncryption(password),
        )
    )


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="document-workbench-pdf-test-") as temporary:
        root = Path(temporary)
        pdf = root / "完整功能.pdf"
        core.create_blank(pdf, 3)
        core.operate_pdf(pdf, "add_text", {"page": 0, "text": "全能文件工作台 測試文字", "font_size": 14}, None)
        core.operate_pdf(pdf, "add_text", {"page": 1, "text": "原始詞 法律文件 合約", "font_size": 12}, None)
        core.operate_pdf(pdf, "header_footer", {"footer": "第 {page} 頁，共 {pages} 頁", "header": "{date}"}, None)
        assert core.search_pdf(pdf, {"text": "原始詞"})["matches"] == 1
        assert core.operate_pdf(pdf, "edit_text", {"text": "原始詞", "replacement": "替代詞", "font_size": 12}, None)["matches"] == 1
        assert core.search_pdf(pdf, {"text": "替代詞"})["matches"] == 1

        core.operate_pdf(pdf, "note", {"page": 0, "text": "便利貼"}, None)
        core.operate_pdf(pdf, "free_text", {"page": 0, "text": "可編輯文字框", "x": 72, "y": 130}, None)
        assert core.operate_pdf(pdf, "mark_search", {"text": "法律文件", "style": "underline"}, None)["matches"] == 1
        core.operate_pdf(pdf, "shape", {"page": 0, "kind": "rectangle", "x": 70, "y": 180, "width": 120, "height": 45}, None)
        core.operate_pdf(pdf, "shape", {"page": 0, "kind": "arrow", "x": 70, "y": 240, "width": 120, "height": 45}, None)
        measurement = core.operate_pdf(pdf, "measure", {"page": 0, "mode": "distance", "x": 70, "y": 290, "width": 30, "height": 40}, None)
        assert round(measurement["measurement"]) == 50
        core.operate_pdf(pdf, "bookmark", {"page": 0, "title": "第一章", "level": 1}, None)
        core.operate_pdf(pdf, "link", {"page": 0, "uri": "https://example.com", "x": 72, "y": 300}, None)

        attachment = root / "附件.txt"
        attachment.write_text("附件內容", encoding="utf-8")
        core.operate_pdf(pdf, "attach_file", {"attachment": str(attachment)}, None)
        attachments = core.attachments_pdf(pdf, {})
        assert attachments["count"] == 1
        extracted_attachment = root / "擷取附件.txt"
        core.operate_pdf(pdf, "extract_attachment", {"name": attachments["attachments"][0]["name"]}, extracted_attachment)
        assert extracted_attachment.read_text(encoding="utf-8") == "附件內容"
        report = core.inspect_pdf(pdf)
        assert report["annotations"] >= 4
        assert report["bookmarks"] == 1
        assert report["links"] == 1
        assert report["attachments"] == 1
        assert core.annotations_pdf(pdf, {})["count"] >= 5

        core.operate_pdf(pdf, "create_field", {"page": 0, "field_type": "text", "name": "姓名", "x": 72, "y": 360}, None)
        core.operate_pdf(pdf, "create_field", {"page": 0, "field_type": "checkbox", "name": "同意", "x": 72, "y": 410, "width": 24, "height": 24}, None)
        core.operate_pdf(pdf, "create_field", {"page": 0, "field_type": "choice", "name": "類別", "choices": "甲,乙", "x": 72, "y": 450}, None)
        fields = core.form_fields_pdf(pdf, {})
        assert fields["count"] == 3
        filled = core.operate_pdf(pdf, "fill_form", {"values": {"姓名": "王小明", "同意": "Yes", "類別": "乙"}}, None)
        assert filled["fields_changed"] == 3
        form_json = root / "表單資料.json"
        core.operate_pdf(pdf, "export_form_data", {}, form_json)
        assert "王小明" in form_json.read_text(encoding="utf-8")
        form_json.write_text('{"姓名":"陳美玲","同意":"Off","類別":"甲"}', encoding="utf-8")
        imported = core.operate_pdf(pdf, "import_form_data", {"data": str(form_json)}, None)
        assert imported["fields_changed"] == 3

        image_a = root / "image-a.png"
        image_b = root / "image-b.png"
        make_image(image_a, 0xCC3333)
        make_image(image_b, 0x3366CC)
        image_pdf = root / "圖片編輯.pdf"
        document = fitz.open()
        page = document.new_page()
        page.insert_image(fitz.Rect(72, 72, 220, 180), filename=str(image_a))
        document.save(image_pdf)
        document.close()
        changed = core.operate_pdf(image_pdf, "image_replace", {"page": 0, "image_index": 0, "image": str(image_b)}, None)
        assert changed["images_changed"] == 1
        deleted = core.operate_pdf(image_pdf, "image_delete", {"page": 0, "image_index": 0}, None)
        assert deleted["images_changed"] == 1

        audit = core.audit_pdf(pdf, {})
        assert audit["pages"] == 3 and "issues" in audit
        core.operate_pdf(pdf, "accessibility_metadata", {"title": "完整 PDF 測試", "author": "全能文件工作台", "language": "zh-TW"}, None)
        assert core.audit_pdf(pdf, {})["metadata"]["title"] == "完整 PDF 測試"

        exports = {
            "txt": root / "匯出.txt",
            "html": root / "匯出.html",
            "png": root / "匯出.png",
            "docx": root / "匯出.docx",
            "xlsx": root / "匯出.xlsx",
            "pptx": root / "匯出.pptx",
        }
        for export_format, destination in exports.items():
            core.operate_pdf(pdf, "export", {"format": export_format, "page": 0}, destination)
            assert destination.is_file() and destination.stat().st_size > 0
        assert "替代詞" in exports["txt"].read_text(encoding="utf-8")
        assert len(WordDocument(exports["docx"]).paragraphs) > 0
        assert len(load_workbook(exports["xlsx"]).worksheets) == 3
        assert len(Presentation(exports["pptx"]).slides) == 3

        encrypted = root / "加密.pdf"
        decrypted = root / "解密.pdf"
        core.operate_pdf(pdf, "encrypt", {"owner_password": "owner", "user_password": "reader"}, encrypted)
        assert core.inspect_pdf(encrypted)["locked"] is True
        core.operate_pdf(encrypted, "decrypt", {"password": "reader"}, decrypted)
        assert core.inspect_pdf(decrypted)["encrypted"] is False

        certificate = root / "test.p12"
        signed = root / "已簽署.pdf"
        make_certificate(certificate, b"secret")
        core.operate_pdf(pdf, "sign", {"certificate": str(certificate), "certificate_password": "secret", "field_name": "Signature1"}, signed)
        signatures = core.verify_signatures_pdf(signed, {})
        assert signatures["available"] and signatures["count"] == 1
        assert signatures["signatures"][0]["intact"] is True

        redacted = root / "永久遮蔽.pdf"
        shutil.copy2(pdf, redacted)
        assert core.operate_pdf(redacted, "redact_search", {"text": "合約"}, None)["matches"] == 1
        assert core.search_pdf(redacted, {"text": "合約"})["matches"] == 0

        pattern_pdf = root / "個資遮蔽.pdf"
        core.create_blank(pattern_pdf, 1)
        core.operate_pdf(pattern_pdf, "add_text", {"text": "身分證 A123456789，信箱 user@example.com"}, None)
        assert core.operate_pdf(pattern_pdf, "redact_pattern", {"pattern": r"[A-Z][12]\d{8}"}, None)["matches"] == 1
        assert core.search_pdf(pattern_pdf, {"text": "A123456789"})["matches"] == 0

        layered = root / "圖層.pdf"
        layered_doc = fitz.open()
        layered_page = layered_doc.new_page()
        ocg = layered_doc.add_ocg("測試圖層", on=True)
        layered_page.insert_text((72, 72), "Layer", oc=ocg)
        layered_doc.save(layered)
        layered_doc.close()
        layers = core.layers_pdf(layered, {})
        assert layers["count"] >= 1
        core.operate_pdf(layered, "layer_visibility", {"number": layers["layers"][0]["number"]}, None)

        filed = core.operate_pdf(pdf, "smart_file", {}, root / "歸檔")
        assert filed["category"] == "合約" and Path(filed["output"]).is_file()

        page_copy = root / "頁面操作.pdf"
        shutil.copy2(pdf, page_copy)
        core.operate_pdf(page_copy, "rotate", {"pages": [0], "angle": 90}, None)
        core.operate_pdf(page_copy, "insert_blank", {"position": 1}, None)
        core.operate_pdf(page_copy, "reorder", {"order": [1, 0, 2, 3]}, None)
        extracted = root / "擷取.pdf"
        core.operate_pdf(page_copy, "extract", {"pages": [0, 2]}, extracted)
        assert core.inspect_pdf(extracted)["pages"] == 2
        split = core.operate_pdf(page_copy, "split", {"output_dir": str(root / "分割")}, None)
        assert len(split["outputs"]) == 4
        core.operate_pdf(page_copy, "delete", {"pages": [3]}, None)
        assert core.live_validate_pdf(page_copy)["passed"] is True
        assert core.compare_pdf(pdf, page_copy)["identical"] is False

        flattened = root / "扁平化.pdf"
        shutil.copy2(pdf, flattened)
        core.operate_pdf(flattened, "flatten", {"annotations": True, "forms": True}, None)
        assert core.inspect_pdf(flattened)["annotations"] == 0
        assert core.inspect_pdf(flattened)["form_fields"] == 0

        status = core.status()
        assert status["protocol_version"] == 2 and status["embedded"] is True
        print("內建 AcroPDF 核心完整往返測試：PASS")


if __name__ == "__main__":
    main()
